import os
from tkinter import ttk, StringVar
from tkinter import messagebox, filedialog

import customtkinter as ctk
from tkcalendar import DateEntry
from PIL import Image

from query import (
    get_project_detail,
    get_transactions,
    delete_transaction
)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer
)
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors

from docx import Document


class TransactionPage(ctk.CTkFrame):

    def __init__(self, parent, project_id):
        super().__init__(parent, fg_color="transparent")

        self.project_id = project_id
        self.selected_group_key = None
        self.raw_transactions = []
        self.grouped_transactions = []

        self.pack(
            fill="both",
            expand=True,
            padx=15,
            pady=15
        )

        project = get_project_detail(project_id)

        if project:
            self.nama_proyek = project["nama_proyek"]
            self.nama_client = project["nama_client"]
        else:
            self.nama_proyek = "Proyek Tidak Ditemukan"
            self.nama_client = "-"

        self.build_ui()
        self.load_data()

    def group_transactions_data(self, raw_data):
        groups = {}

        for trx in raw_data:
            tanggal_str = str(trx["tanggal"])
            jenis = trx["jenis"]
            kategori = trx.get("kategori", "-") or "-"

            key = (tanggal_str, jenis)

            if key not in groups:
                groups[key] = {
                    "tanggal": tanggal_str,
                    "jenis": jenis,
                    "items": [],
                    "total_nominal": 0.0,
                    "ids": []
                }

            ket = trx.get("keterangan", "-")
            nom = float(trx["nominal"])
            nota = trx.get("nota", "")

            groups[key]["items"].append({
                "id": trx["id"],
                "kategori": kategori,
                "keterangan": ket,
                "nominal": nom,
                "nota": nota
            })

            groups[key]["total_nominal"] += nom
            groups[key]["ids"].append(trx["id"])

        result = []
        for key, data in groups.items():
            kategori_str = "\n".join([item["kategori"] for item in data["items"]])
            keterangan_str = "\n".join([item["keterangan"] for item in data["items"]])
            nominal_str = "\n".join([f"Rp {item['nominal']:,.0f}" for item in data["items"]])
            
            nota_lines = []
            for item in data["items"]:
                if item["nota"]:
                    nota_lines.append("Lihat Nota")
                else:
                    nota_lines.append("-")
            nota_str = "\n".join(nota_lines)

            result.append({
                "group_key": f"{data['tanggal']}_{data['jenis']}",
                "tanggal": data["tanggal"],
                "kategori_str": kategori_str if kategori_str else "-",
                "keterangan_str": keterangan_str if keterangan_str else "-",
                "nominal_items_str": nominal_str if nominal_str else "-",
                "jenis": data["jenis"],
                "total_nominal": data["total_nominal"],
                "nota_str": nota_str,
                "ids": data["ids"],
                "items": data["items"]
            })

        result.sort(key=lambda x: x["tanggal"], reverse=True)
        return result

    def calculate_totals(self, data_list):
        total_pemasukan = sum(item["total_nominal"] for item in data_list if item["jenis"] == "Pemasukan")
        total_pengeluaran = sum(item["total_nominal"] for item in data_list if item["jenis"] == "Pengeluaran")
        saldo = total_pemasukan - total_pengeluaran

        if saldo == 0:
            status = "Balance (Rp 0)"
        elif saldo > 0:
            status = f"Surplus: Rp {saldo:,.0f}"
        else:
            status = f"Defisit: Rp {abs(saldo):,.0f}"

        return total_pemasukan, total_pengeluaran, saldo, status

    def export_excel(self):
        data_to_export = self.grouped_transactions
        if not data_to_export:
            messagebox.showwarning("Peringatan", "Tidak ada data untuk diexport.")
            return

        file = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel File", "*.xlsx")])
        if not file:
            return

        wb = Workbook()
        ws = wb.active
        ws.title = "Laporan Transaksi"

        ws.append([f"Laporan Transaksi: {self.nama_proyek}"])
        ws["A1"].font = Font(name="Arial", size=14, bold=True)
        ws.append([f"Client: {self.nama_client}"])
        ws["A2"].font = Font(name="Arial", size=11, italic=True)
        ws.append([])

        headers = ["No", "Tanggal", "Kategori", "Detail Barang", "Harga Barang", "Jenis", "Total"]
        ws.append(headers)
        
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=4, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        data_font = Font(name="Arial", size=10)
        for no, item in enumerate(data_to_export, 1):
            ws.append([
                no, item["tanggal"], item["kategori_str"], item["keterangan_str"],
                item["nominal_items_str"], item["jenis"], item["total_nominal"]
            ])
            
            curr_row = ws.max_row
            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=curr_row, column=col_idx)
                c.font = data_font
                c.alignment = Alignment(vertical="top", wrap_text=True)

        total_pemasukan, total_pengeluaran, saldo, status = self.calculate_totals(data_to_export)
        ws.append([])
        ws.append(["", "", "", "", "Total Pemasukan:", total_pemasukan])
        ws.append(["", "", "", "", "Total Pengeluaran:", total_pengeluaran])
        ws.append(["", "", "", "", "Saldo / Selisih:", saldo])
        ws.append(["", "", "", "", "Status Keuangan:", status])

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                lines = val_str.split('\n')
                for line in lines:
                    if len(line) > max_len:
                        max_len = len(line)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        wb.save(file)
        messagebox.showinfo("Berhasil", "Export Excel berhasil")

    def export_pdf(self):
        data_to_export = self.grouped_transactions
        if not data_to_export:
            messagebox.showwarning("Peringatan", "Tidak ada data untuk diexport.")
            return

        file = filedialog.asksaveasfilename(defaultextension=".pdf", filetypes=[("PDF File", "*.pdf")])
        if not file:
            return

        pdf = SimpleDocTemplate(file)
        styles = getSampleStyleSheet()
        story = []

        story.append(Paragraph(f"<b>Laporan Transaksi: {self.nama_proyek}</b>", styles["Heading1"]))
        story.append(Paragraph(f"Client: {self.nama_client}", styles["Normal"]))
        story.append(Spacer(1, 15))

        table_data = [["No", "Tanggal", "Kategori", "Detail Barang", "Harga Barang", "Jenis", "Total"]]
        for no, item in enumerate(data_to_export, 1):
            table_data.append([
                str(no), item["tanggal"],
                Paragraph(item["kategori_str"].replace("\n", "<br/>"), styles["Normal"]),
                Paragraph(item["keterangan_str"].replace("\n", "<br/>"), styles["Normal"]),
                Paragraph(item["nominal_items_str"].replace("\n", "<br/>"), styles["Normal"]),
                item["jenis"], f"Rp {item['total_nominal']:,.0f}"
            ])

        table = Table(table_data, colWidths=[20, 60, 80, 130, 85, 55, 75])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563eb")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey)
        ]))
        story.append(table)
        story.append(Spacer(1, 15))

        total_pemasukan, total_pengeluaran, saldo, status = self.calculate_totals(data_to_export)
        summary_data = [
            ["Total Pemasukan", f"Rp {total_pemasukan:,.0f}"],
            ["Total Pengeluaran", f"Rp {total_pengeluaran:,.0f}"],
            ["Saldo / Selisih", f"Rp {saldo:,.0f}"],
            ["Status Keuangan", status]
        ]
        summary_table = Table(summary_data, colWidths=[150, 200])
        summary_table.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f3f4f6"))
        ]))
        story.append(summary_table)
        pdf.build(story)
        messagebox.showinfo("Berhasil", "Export PDF berhasil")

    def export_word(self):
        data_to_export = self.grouped_transactions
        if not data_to_export:
            messagebox.showwarning("Peringatan", "Tidak ada data untuk diexport.")
            return

        file = filedialog.asksaveasfilename(defaultextension=".docx", filetypes=[("Word File", "*.docx")])
        if not file:
            return

        doc = Document()
        doc.add_heading(f"Laporan Transaksi: {self.nama_proyek}", level=1)
        doc.add_paragraph(f"Client: {self.nama_client}")

        table = doc.add_table(rows=1, cols=7)
        table.style = "Table Grid"
        header = table.rows[0].cells
        for idx, text in enumerate(["No", "Tanggal", "Kategori", "Detail Barang", "Harga Barang", "Jenis", "Total"]):
            header[idx].text = text

        for no, item in enumerate(data_to_export, 1):
            row = table.add_row().cells
            row[0].text = str(no)
            row[1].text = item["tanggal"]
            row[2].text = item["kategori_str"]
            row[3].text = item["keterangan_str"]
            row[4].text = item["nominal_items_str"]
            row[5].text = item["jenis"]
            row[6].text = f"Rp {item['total_nominal']:,.0f}"

        total_pemasukan, total_pengeluaran, saldo, status = self.calculate_totals(data_to_export)
        doc.add_paragraph()
        p_sum = doc.add_paragraph()
        p_sum.add_run("RINGKASAN LAPORAN:\n").bold = True
        p_sum.add_run(f"• Total Pemasukan : Rp {total_pemasukan:,.0f}\n")
        p_sum.add_run(f"• Total Pengeluaran : Rp {total_pengeluaran:,.0f}\n")
        p_sum.add_run(f"• Saldo / Selisih : Rp {saldo:,.0f}\n")
        p_sum.add_run(f"• Status Keuangan : {status}\n").bold = True

        doc.save(file)
        messagebox.showinfo("Berhasil", "Export Word berhasil")

    def show_nota_popup(self, items):
        valid_items = [i for i in items if i.get("nota") and os.path.exists(i["nota"])]
        if not valid_items:
            messagebox.showinfo("Info Nota", "Tidak ada file foto nota yang tersedia.")
            return

        top = ctk.CTkToplevel(self)
        top.title("Preview Nota Transaksi")
        top.geometry("560x700")
        top.configure(fg_color=("#f8fafc", "#0f172a"))
        top.grab_set()

        ctk.CTkLabel(
            top, 
            text="🧾 Detail & Pratinjau Nota Transaksi", 
            font=("Arial", 16, "bold"),
            text_color="#7c3aed"
        ).pack(pady=(15, 5))

        scroll_frame = ctk.CTkScrollableFrame(top, fg_color=("#ffffff", "#0f172a"))
        scroll_frame.pack(fill="both", expand=True, padx=15, pady=10)

        for item in valid_items:
            card = ctk.CTkFrame(
                scroll_frame, 
                fg_color=("#f8fafc", "#1e293b"), 
                corner_radius=10, 
                border_width=1, 
                border_color=("#e2e8f0", "#334155")
            )
            card.pack(fill="x", pady=8, padx=5)

            ctk.CTkLabel(
                card, 
                text=f"📌 [{item['kategori']}] {item['keterangan']} - Rp {item['nominal']:,.0f}", 
                font=("Arial", 12, "bold"),
                text_color=("#0f172a", "#ffffff")
            ).pack(anchor="w", padx=12, pady=(10, 6))

            try:
                pil_img = Image.open(item["nota"])
                
                max_w, max_h = 480, 450
                img_w, img_h = pil_img.size
                ratio = min(max_w / img_w, max_h / img_h)
                new_size = (int(img_w * ratio), int(img_h * ratio))

                ctk_img = ctk.CTkImage(light_image=pil_img, dark_image=pil_img, size=new_size)
                img_label = ctk.CTkLabel(card, image=ctk_img, text="")
                img_label.pack(pady=(4, 10), padx=10)

            except Exception as e:
                ctk.CTkLabel(card, text=f"Gagal memuat foto: {e}", text_color="#ef4444").pack(pady=10)

    def update_summary(self, total_pemasukan, total_pengeluaran):
        saldo = total_pemasukan - total_pengeluaran

        self.lbl_pemasukan_bottom.configure(text=f"Rp {total_pemasukan:,.0f}")
        self.lbl_pengeluaran_bottom.configure(text=f"Rp {total_pengeluaran:,.0f}")
        self.lbl_saldo_bottom.configure(text=f"Rp {saldo:,.0f}")

        if saldo == 0:
            status_text = "BALANCE"
            color = "#3b82f6"
        elif saldo > 0:
            status_text = f"SURPLUS\n(+Rp {saldo:,.0f})"
            color = "#10b981"
        else:
            status_text = f"DEFISIT\n(-Rp {abs(saldo):,.0f})"
            color = "#ef4444"

        self.lbl_balance_status.configure(text=status_text, text_color=color)
        self.lbl_saldo.configure(text=f"Saldo : Rp {saldo:,.0f}")

    def load_data(self):
        for item in self.tree.get_children():
            self.tree.delete(item)

        transactions = get_transactions(self.project_id)
        self.raw_transactions = transactions
        self.grouped_transactions = self.group_transactions_data(transactions)

        # -------------------------------------------------------------
        # KONFIGURASI STYLE TABEL REAL-TIME SESUAI MODE TAMPILAN (DARK/LIGHT)
        # -------------------------------------------------------------
        is_dark = ctk.get_appearance_mode().lower() == "dark"
        
        table_bg = "#0f172a" if is_dark else "#ffffff"
        text_fg = "#ffffff" if is_dark else "#0f172a"
        header_bg = "#1e293b" if is_dark else "#f1f5f9"
        pemasukan_bg = "#064e3b" if is_dark else "#d1fae5"
        pengeluaran_bg = "#7f1d1d" if is_dark else "#fee2e2"

        style = ttk.Style()
        style.theme_use("clam")
        
        # Style Header & Isi Tabel
        style.configure("Treeview.Heading", background=header_bg, foreground=text_fg, font=("Arial", 10, "bold"))
        style.configure("Treeview", background=table_bg, foreground=text_fg, fieldbackground=table_bg, rowheight=55, font=("Arial", 10))

        # Tag Warna Baris Pemasukan & Pengeluaran
        self.tree.tag_configure("pemasukan", background=pemasukan_bg, foreground=text_fg)
        self.tree.tag_configure("pengeluaran", background=pengeluaran_bg, foreground=text_fg)

        total_pemasukan = 0
        total_pengeluaran = 0
        current_date = None
        nomor_display = 1

        for item in self.grouped_transactions:
            if item["jenis"] == "Pemasukan":
                total_pemasukan += item["total_nominal"]
                warna = "pemasukan"
            else:
                total_pengeluaran += item["total_nominal"]
                warna = "pengeluaran"

            if current_date != item["tanggal"]:
                current_date = item["tanggal"]
                no_col = str(nomor_display)
                nomor_display += 1
            else:
                no_col = ""

            self.tree.insert(
                "",
                "end",
                iid=item["group_key"],
                values=(
                    no_col,
                    item["tanggal"],
                    item["kategori_str"],
                    item["keterangan_str"],
                    item["nominal_items_str"],
                    item["jenis"],
                    f"Rp {item['total_nominal']:,.0f}",
                    item["nota_str"]
                ),
                tags=(warna,)
            )

        self.update_summary(total_pemasukan, total_pengeluaran)

    def on_select(self, event):
        selected = self.tree.selection()
        self.selected_group_key = selected[0] if selected else None

    def on_double_click(self, event):
        item_id = self.tree.identify_row(event.y)
        column = self.tree.identify_column(event.x)
        if item_id and column == "#8":
            matched = next((x for x in self.grouped_transactions if x["group_key"] == item_id), None)
            if matched:
                self.show_nota_popup(matched["items"])

    def open_add_transaction(self):
        from views.add_transaction import AddTransaction
        AddTransaction(self, self.project_id, self.load_data)

    def edit_transaction(self):
        if not self.selected_group_key:
            messagebox.showwarning("Pilih Data", "Silahkan pilih baris transaksi terlebih dahulu")
            return
            
        matched = next((x for x in self.grouped_transactions if x["group_key"] == self.selected_group_key), None)
        if not matched:
            return

        if len(matched["items"]) == 1:
            from views.edit_transaction import EditTransaction
            EditTransaction(self, matched["items"][0]["id"], self.load_data)
        else:
            self.show_edit_item_popup(matched)

    def show_edit_item_popup(self, matched_group):
        popup = ctk.CTkToplevel(self)
        popup.title("Pilih Detail Barang yang Ingin Di-edit")
        popup.geometry("420x360")
        popup.configure(fg_color=("#f8fafc", "#0f172a"))
        popup.grab_set()

        ctk.CTkLabel(
            popup, 
            text="Pilih item transaksi yang ingin diubah:", 
            font=("Arial", 13, "bold"),
            text_color="#7c3aed"
        ).pack(pady=(15, 5))

        frame_list = ctk.CTkScrollableFrame(popup, fg_color=("#ffffff", "#1e293b"), corner_radius=10)
        frame_list.pack(pady=10, padx=15, fill="both", expand=True)

        selected_id_var = ctk.IntVar(value=matched_group["items"][0]["id"])
        for item in matched_group["items"]:
            rb = ctk.CTkRadioButton(
                frame_list,
                text=f"[{item['kategori']}] {item['keterangan']} - Rp {item['nominal']:,.0f}",
                value=item["id"],
                variable=selected_id_var,
                text_color=("#0f172a", "#ffffff")
            )
            rb.pack(anchor="w", pady=6, padx=8)

        def do_edit():
            target_id = selected_id_var.get()
            popup.destroy()
            from views.edit_transaction import EditTransaction
            EditTransaction(self, target_id, self.load_data)

        btn_frame = ctk.CTkFrame(popup, fg_color="transparent")
        btn_frame.pack(fill="x", pady=(0, 15), padx=15)

        ctk.CTkButton(
            btn_frame, 
            text="Lanjut Edit", 
            fg_color="#f59e0b", 
            hover_color="#d97706", 
            command=do_edit
        ).pack(side="right", padx=5)

        ctk.CTkButton(
            btn_frame, 
            text="Batal", 
            fg_color="#64748b", 
            hover_color="#475569", 
            command=popup.destroy
        ).pack(side="right", padx=5)

    def delete_transaction(self):
        if not self.selected_group_key:
            messagebox.showwarning("Pilih Data", "Silahkan pilih transaksi yang ingin dihapus")
            return
        matched = next((x for x in self.grouped_transactions if x["group_key"] == self.selected_group_key), None)
        if not matched:
            return

        if len(matched["items"]) == 1:
            target_item = matched["items"][0]
            if messagebox.askyesno("Konfirmasi Hapus", f"Apakah Anda yakin ingin menghapus transaksi '{target_item['keterangan']}'?"):
                delete_transaction(target_item["id"])
                messagebox.showinfo("Berhasil", "Transaksi berhasil dihapus")
                self.selected_group_key = None
                self.load_data()
        else:
            self.show_delete_item_popup(matched)

    def show_delete_item_popup(self, matched_group):
        popup = ctk.CTkToplevel(self)
        popup.title("Pilih Detail Barang yang Ingin Dihapus")
        popup.geometry("400x350")
        popup.configure(fg_color=("#f8fafc", "#0f172a"))
        popup.grab_set()

        ctk.CTkLabel(popup, text="Pilih item transaksi:", font=("Arial", 12, "bold"), text_color=("#0f172a", "#ffffff")).pack(pady=10)
        frame_list = ctk.CTkScrollableFrame(popup, width=360, height=200, fg_color=("#ffffff", "#1e293b"))
        frame_list.pack(pady=5, padx=10, fill="both", expand=True)

        selected_id_var = ctk.IntVar(value=matched_group["items"][0]["id"])
        for item in matched_group["items"]:
            rb = ctk.CTkRadioButton(
                frame_list,
                text=f"[{item['kategori']}] {item['keterangan']} (Rp {item['nominal']:,.0f})",
                value=item["id"],
                variable=selected_id_var,
                text_color=("#0f172a", "#ffffff")
            )
            rb.pack(anchor="w", pady=5, padx=5)

        def do_delete():
            target_id = selected_id_var.get()
            item_to_delete = next((i for i in matched_group["items"] if i["id"] == target_id), None)
            item_name = item_to_delete["keterangan"] if item_to_delete else ""

            if messagebox.askyesno("Konfirmasi Hapus", f"Apakah Anda yakin ingin menghapus item '{item_name}'?"):
                delete_transaction(target_id)
                popup.destroy()
                messagebox.showinfo("Berhasil", "Item transaksi berhasil dihapus")
                self.selected_group_key = None
                self.load_data()

        btn_frame = ctk.CTkFrame(popup, fg_color="transparent")
        btn_frame.pack(fill="x", pady=10)
        ctk.CTkButton(btn_frame, text="Hapus", fg_color="#dc2626", command=do_delete).pack(side="right", padx=10)
        ctk.CTkButton(btn_frame, text="Batal", fg_color="#6b7280", command=popup.destroy).pack(side="right", padx=5)

    def refresh_data(self):
        self.search.delete(0, "end")
        self.var_start.set("")
        self.var_end.set("")
        self.load_data()

    def search_data(self, event=None):
        keyword = self.search.get().lower().strip()
        
        start_date = self.var_start.get().strip()
        end_date = self.var_end.get().strip()

        for item in self.tree.get_children():
            self.tree.delete(item)

        filtered_raw = []
        for trx in self.raw_transactions:
            trx_date = str(trx["tanggal"])

            if start_date and trx_date < start_date:
                continue
            if end_date and trx_date > end_date:
                continue

            searchable_text = f"{trx_date} {trx.get('jenis','')} {trx.get('kategori','')} {trx.get('keterangan','')} {trx.get('nominal','')}".lower()
            if keyword in searchable_text:
                filtered_raw.append(trx)

        display_groups = self.group_transactions_data(filtered_raw)
        total_pemasukan = 0
        total_pengeluaran = 0
        current_date = None
        nomor_display = 1

        for item in display_groups:
            if item["jenis"] == "Pemasukan":
                total_pemasukan += item["total_nominal"]
                warna = "pemasukan"
            else:
                total_pengeluaran += item["total_nominal"]
                warna = "pengeluaran"

            if current_date != item["tanggal"]:
                current_date = item["tanggal"]
                no_col = str(nomor_display)
                nomor_display += 1
            else:
                no_col = ""

            self.tree.insert(
                "", "end", iid=item["group_key"],
                values=(no_col, item["tanggal"], item["kategori_str"], item["keterangan_str"], item["nominal_items_str"], item["jenis"], f"Rp {item['total_nominal']:,.0f}", item["nota_str"]),
                tags=(warna,)
            )

        self.update_summary(total_pemasukan, total_pengeluaran)

    def build_ui(self):
        # 1. Header
        header = ctk.CTkFrame(self, fg_color="transparent")
        header.pack(fill="x", pady=(0, 5))

        ctk.CTkLabel(header, text=self.nama_proyek, font=("Arial", 24, "bold"), text_color=("#0f172a", "#ffffff")).pack(anchor="w")
        ctk.CTkLabel(header, text=f"Client : {self.nama_client}", font=("Arial", 14), text_color=("#64748b", "#94a3b8")).pack(anchor="w")
        self.lbl_saldo = ctk.CTkLabel(header, text="Saldo : Rp 0", font=("Arial", 16, "bold"), text_color="#10b981")
        self.lbl_saldo.pack(anchor="w", pady=2)

        # 2. Toolbar
        toolbar = ctk.CTkFrame(self, fg_color=("#ffffff", "#1e293b"), corner_radius=12, border_width=1, border_color=("#e2e8f0", "#334155"))
        toolbar.pack(fill="x", pady=5)

        button_frame = ctk.CTkFrame(toolbar, fg_color="transparent")
        button_frame.pack(side="left", padx=5, pady=5)

        ctk.CTkButton(button_frame, text="+ Tambah", width=85, fg_color="#7c3aed", hover_color="#6d28d9", command=self.open_add_transaction).pack(side="left", padx=2)
        ctk.CTkButton(button_frame, text="✏ Edit", width=75, fg_color="#f59e0b", hover_color="#d97706", command=self.edit_transaction).pack(side="left", padx=2)
        ctk.CTkButton(button_frame, text="🗑 Hapus", width=75, fg_color="#dc2626", hover_color="#b91c1c", command=self.delete_transaction).pack(side="left", padx=2)
        ctk.CTkButton(button_frame, text="↻ Refresh", width=80, command=self.refresh_data).pack(side="left", padx=2)
        ctk.CTkButton(button_frame, text="📊 Excel", width=75, fg_color="#16a34a", hover_color="#15803d", command=self.export_excel).pack(side="left", padx=2)
        ctk.CTkButton(button_frame, text="📄 PDF", width=75, fg_color="#dc2626", hover_color="#b91c1c", command=self.export_pdf).pack(side="left", padx=2)
        ctk.CTkButton(button_frame, text="📝 Word", width=75, fg_color="#2563eb", hover_color="#1d4ed8", command=self.export_word).pack(side="left", padx=2)

        # 3. Filter Bar
        filter_bar = ctk.CTkFrame(self, fg_color=("#ffffff", "#1e293b"), corner_radius=12, border_width=1, border_color=("#e2e8f0", "#334155"))
        filter_bar.pack(fill="x", pady=(0, 5))

        self.var_start = StringVar(value="")
        self.var_end = StringVar(value="")

        ctk.CTkLabel(filter_bar, text="Dari:", font=("Arial", 11), text_color=("#0f172a", "#ffffff")).pack(side="left", padx=(10, 2))
        self.date_start = DateEntry(
            filter_bar, 
            width=11, 
            background='darkblue', 
            foreground='white', 
            date_pattern='yyyy-mm-dd',
            textvariable=self.var_start
        )
        self.date_start.pack(side="left", padx=3, pady=5)

        ctk.CTkLabel(filter_bar, text="s/d", font=("Arial", 11), text_color=("#0f172a", "#ffffff")).pack(side="left", padx=2)
        self.date_end = DateEntry(
            filter_bar, 
            width=11, 
            background='darkblue', 
            foreground='white', 
            date_pattern='yyyy-mm-dd',
            textvariable=self.var_end
        )
        self.date_end.pack(side="left", padx=3, pady=5)

        ctk.CTkButton(filter_bar, text="Filter Tgl", width=75, fg_color="#7c3aed", command=self.search_data).pack(side="left", padx=5)

        self.search = ctk.CTkEntry(filter_bar, width=200, placeholder_text="Cari transaksi...")
        self.search.pack(side="right", padx=10, pady=5)
        self.search.bind("<KeyRelease>", self.search_data)

        # 4. Bottom Summary Panel
        summary_frame = ctk.CTkFrame(self, fg_color="transparent")
        summary_frame.pack(side="bottom", fill="x", pady=(8, 0))

        summary_frame.columnconfigure((0, 1, 2, 3), weight=1, uniform="card")

        c1 = ctk.CTkFrame(summary_frame, fg_color=("#dcfce7", "#064e3b"), corner_radius=8)
        c1.grid(row=0, column=0, padx=4, pady=2, sticky="ew")
        ctk.CTkLabel(c1, text="Total Pemasukan", font=("Arial", 10, "bold"), text_color=("#166534", "#a7f3d0")).pack(pady=(6, 0))
        self.lbl_pemasukan_bottom = ctk.CTkLabel(c1, text="Rp 0", font=("Arial", 13, "bold"), text_color=("#0f172a", "#ffffff"))
        self.lbl_pemasukan_bottom.pack(pady=(0, 6))

        c2 = ctk.CTkFrame(summary_frame, fg_color=("#fee2e2", "#7f1d1d"), corner_radius=8)
        c2.grid(row=0, column=1, padx=4, pady=2, sticky="ew")
        ctk.CTkLabel(c2, text="Total Pengeluaran", font=("Arial", 10, "bold"), text_color=("#991b1b", "#fecaca")).pack(pady=(6, 0))
        self.lbl_pengeluaran_bottom = ctk.CTkLabel(c2, text="Rp 0", font=("Arial", 13, "bold"), text_color=("#0f172a", "#ffffff"))
        self.lbl_pengeluaran_bottom.pack(pady=(0, 6))

        c3 = ctk.CTkFrame(summary_frame, fg_color=("#dbeafe", "#1e3a8a"), corner_radius=8)
        c3.grid(row=0, column=2, padx=4, pady=2, sticky="ew")
        ctk.CTkLabel(c3, text="Total Saldo", font=("Arial", 10, "bold"), text_color=("#1e40af", "#bfdbfe")).pack(pady=(6, 0))
        self.lbl_saldo_bottom = ctk.CTkLabel(c3, text="Rp 0", font=("Arial", 13, "bold"), text_color=("#0f172a", "#ffffff"))
        self.lbl_saldo_bottom.pack(pady=(0, 6))

        c4 = ctk.CTkFrame(summary_frame, fg_color=("#f3f4f6", "#374151"), corner_radius=8)
        c4.grid(row=0, column=3, padx=4, pady=2, sticky="ew")
        ctk.CTkLabel(c4, text="Status Balance", font=("Arial", 10, "bold"), text_color=("#374151", "#d1d5db")).pack(pady=(6, 0))
        self.lbl_balance_status = ctk.CTkLabel(c4, text="BALANCE", font=("Arial", 11, "bold"), text_color="#9ca3af")
        self.lbl_balance_status.pack(pady=(0, 6))

        # 5. Tabel Data (Headers Diubah ke Title Case / Huruf Kapital Awal)
        table_frame = ctk.CTkFrame(self, fg_color="transparent")
        table_frame.pack(side="top", fill="both", expand=True, pady=5)

        style = ttk.Style()
        style.theme_use("clam")
        style.configure("Treeview", rowheight=55, font=("Arial", 10))

        self.tree = ttk.Treeview(
            table_frame,
            columns=("no", "tanggal", "kategori", "keterangan", "nominal_items", "jenis", "nominal", "nota"),
            show="headings"
        )

        headers = [
            ("no", "No", 40, 35),
            ("tanggal", "Tanggal", 90, 85),
            ("kategori", "Kategori", 120, 100),
            ("keterangan", "Detail Barang", 220, 160),
            ("nominal_items", "Harga Barang", 120, 100),
            ("jenis", "Jenis", 85, 75),
            ("nominal", "Total", 120, 100),
            ("nota", "Nota", 100, 90)
        ]

        for col, text, width, min_w in headers:
            self.tree.heading(col, text=text)
            self.tree.column(
                col, 
                width=width, 
                minwidth=min_w, 
                stretch=True, 
                anchor="center" if col not in ["kategori", "keterangan", "nominal_items", "nota"] else "w"
            )

        v_scrollbar = ttk.Scrollbar(table_frame, orient="vertical", command=self.tree.yview)
        h_scrollbar = ttk.Scrollbar(table_frame, orient="horizontal", command=self.tree.xview)

        self.tree.configure(yscrollcommand=v_scrollbar.set, xscrollcommand=h_scrollbar.set)

        self.tree.grid(row=0, column=0, sticky="nsew")
        v_scrollbar.grid(row=0, column=1, sticky="ns")
        h_scrollbar.grid(row=1, column=0, sticky="ew")

        table_frame.grid_rowconfigure(0, weight=1)
        table_frame.grid_columnconfigure(0, weight=1)

        self.tree.bind("<<TreeviewSelect>>", self.on_select)
        self.tree.bind("<Double-1>", self.on_double_click)